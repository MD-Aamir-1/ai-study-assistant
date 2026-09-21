"""
File text extraction service.

Strategy:
  1. Try PyMuPDF (fastest, best for text PDFs)
  2. Try pdfplumber
  3. Try pypdf
  4. If all give < 500 chars -> fall back to OCR (PyMuPDF render + Tesseract)

Supports: PDF, DOCX, TXT, MD, CSV, JSON, code files, images.
"""

import io
import os
import re
import shutil
from typing import Dict, Any

# ---------- Optional imports ----------
try:
    import pymupdf
    HAS_PYMUPDF = True
except ImportError:
    try:
        import fitz as pymupdf
        HAS_PYMUPDF = True
    except ImportError:
        HAS_PYMUPDF = False

try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    from pypdf import PdfReader
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import pytesseract
    HAS_TESSERACT = True
except ImportError:
    HAS_TESSERACT = False


# ---------- Auto-detect Tesseract on Windows ----------
def _configure_tesseract():
    if not HAS_TESSERACT:
        return False
    if shutil.which("tesseract"):
        return True
    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.expanduser(r"~\AppData\Local\Tesseract-OCR\tesseract.exe"),
    ]
    for path in candidates:
        if os.path.exists(path):
            pytesseract.pytesseract.tesseract_cmd = path
            return True
    return False


_TESSERACT_READY = _configure_tesseract()
MIN_GOOD_CHARS = 500


# ==================================================
# PDF: PyMuPDF — try multiple text modes
# ==================================================
def _extract_pdf_pymupdf(contents: bytes) -> Dict[str, Any]:
    doc = pymupdf.open(stream=contents, filetype="pdf")
    pages_text = []

    for page in doc:
        best = ""
        try:
            t = page.get_text("text") or ""
            if len(t) > len(best):
                best = t
        except Exception:
            pass
        try:
            blocks = page.get_text("blocks") or []
            t = "\n".join(
                b[4] for b in blocks
                if len(b) >= 5 and isinstance(b[4], str)
            )
            if len(t) > len(best):
                best = t
        except Exception:
            pass
        try:
            d = page.get_text("dict") or {}
            spans = []
            for block in d.get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        s = span.get("text", "")
                        if s:
                            spans.append(s)
            t = " ".join(spans)
            if len(t) > len(best):
                best = t
        except Exception:
            pass

        pages_text.append(best.strip())

    doc.close()

    full = "\n\n".join(
        f"===== Page {i + 1} =====\n{t}" for i, t in enumerate(pages_text)
    )
    return {"pages": len(pages_text), "text": full, "method": "pymupdf"}


# ==================================================
# PDF: pdfplumber
# ==================================================
def _extract_pdf_pdfplumber(contents: bytes) -> Dict[str, Any]:
    pages_text = []
    with pdfplumber.open(io.BytesIO(contents)) as pdf:
        for page in pdf.pages:
            try:
                t = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            except Exception:
                t = ""
            pages_text.append(t.strip())
    full = "\n\n".join(
        f"===== Page {i + 1} =====\n{t}" for i, t in enumerate(pages_text)
    )
    return {"pages": len(pages_text), "text": full, "method": "pdfplumber"}


# ==================================================
# PDF: pypdf
# ==================================================
def _extract_pdf_pypdf(contents: bytes) -> Dict[str, Any]:
    reader = PdfReader(io.BytesIO(contents))
    pages_text = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        pages_text.append(t.strip())
    full = "\n\n".join(
        f"===== Page {i + 1} =====\n{t}" for i, t in enumerate(pages_text)
    )
    return {"pages": len(pages_text), "text": full, "method": "pypdf"}


# ==================================================
# PDF: OCR via PyMuPDF render + Tesseract
# ==================================================
def _extract_pdf_ocr(contents: bytes, max_pages: int = 50) -> Dict[str, Any]:
    if not _TESSERACT_READY:
        raise RuntimeError(
            "PDF appears to be image-based, but Tesseract OCR is not installed. "
            "Install: https://github.com/UB-Mannheim/tesseract/wiki"
        )
    if not HAS_PYMUPDF:
        raise RuntimeError("OCR needs PyMuPDF for rendering. Run: pip install pymupdf")
    if not HAS_PIL:
        raise RuntimeError("OCR needs Pillow. Run: pip install Pillow")

    doc = pymupdf.open(stream=contents, filetype="pdf")
    pages_text = []
    total = len(doc)

    for i, page in enumerate(doc):
        if i >= max_pages:
            break
        try:
            pix = page.get_pixmap(dpi=200)
            img_bytes = pix.tobytes("png")
            img = Image.open(io.BytesIO(img_bytes))
            text = pytesseract.image_to_string(img) or ""
        except Exception:
            text = ""
        pages_text.append(text.strip())

    doc.close()

    full = "\n\n".join(
        f"===== Page {i + 1} =====\n{t}" for i, t in enumerate(pages_text)
    )
    note = ""
    if total > max_pages:
        note = f"OCR'd first {max_pages} of {total} pages"

    return {
        "pages": total,
        "text": full,
        "method": "pymupdf-ocr",
        "note": note,
    }


# ==================================================
# PDF dispatcher
# ==================================================
def _extract_pdf(contents: bytes) -> Dict[str, Any]:
    results = []
    errors = []

    if HAS_PYMUPDF:
        try:
            results.append(_extract_pdf_pymupdf(contents))
        except Exception as e:
            errors.append(f"pymupdf: {e}")

    if HAS_PDFPLUMBER:
        try:
            results.append(_extract_pdf_pdfplumber(contents))
        except Exception as e:
            errors.append(f"pdfplumber: {e}")

    if HAS_PYPDF:
        try:
            results.append(_extract_pdf_pypdf(contents))
        except Exception as e:
            errors.append(f"pypdf: {e}")

    best_text_layer = (
        max(results, key=lambda r: len(r["text"])) if results else None
    )

    if best_text_layer and len(best_text_layer["text"]) >= MIN_GOOD_CHARS:
        return best_text_layer

    try:
        ocr_result = _extract_pdf_ocr(contents)
        results.append(ocr_result)
    except Exception as e:
        errors.append(f"ocr: {e}")

    if results:
        best = max(results, key=lambda r: len(r["text"]))
        if len(best["text"]) < MIN_GOOD_CHARS:
            best["warning"] = (
                "Extraction returned very little text. "
                "The PDF may be protected, empty, or in an unsupported format."
            )
        best["errors"] = errors
        return best

    raise RuntimeError("Could not extract PDF. " + " | ".join(errors))


# ==================================================
# DOCX
# ==================================================
def _extract_docx(contents: bytes) -> Dict[str, Any]:
    if not HAS_DOCX:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")
    doc = Document(io.BytesIO(contents))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return {"pages": 1, "text": "\n".join(parts), "method": "python-docx"}


# ==================================================
# Plain text
# ==================================================
def _extract_plaintext(contents: bytes) -> Dict[str, Any]:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return {"pages": 1, "text": contents.decode(enc), "method": "plain"}
        except UnicodeDecodeError:
            continue
    raise RuntimeError("Could not decode text file.")


# ==================================================
# Image OCR
# ==================================================
def _extract_image(contents: bytes) -> Dict[str, Any]:
    if not (_TESSERACT_READY and HAS_PIL):
        raise RuntimeError(
            "Image OCR needs Tesseract. "
            "Install: https://github.com/UB-Mannheim/tesseract/wiki"
        )
    image = Image.open(io.BytesIO(contents))
    text = pytesseract.image_to_string(image)
    return {"pages": 1, "text": text.strip(), "method": "tesseract-image"}


# ==================================================
# Public API
# ==================================================
def extract_text(contents: bytes, content_type: str, filename: str) -> Dict[str, Any]:
    if not contents:
        raise ValueError("Empty file.")

    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    if name.endswith(".pdf") or "pdf" in ctype:
        r = _extract_pdf(contents)
        r["source_type"] = "pdf"
        return _finalize(r)

    if name.endswith(".docx") or "wordprocessingml" in ctype:
        r = _extract_docx(contents)
        r["source_type"] = "docx"
        return _finalize(r)

    if (
        name.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"))
        or ctype.startswith("image/")
    ):
        r = _extract_image(contents)
        r["source_type"] = "image"
        return _finalize(r)

    if (
        name.endswith((
            ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml",
            ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp",
            ".cs", ".go", ".rs", ".rb", ".php", ".html", ".css", ".sql",
            ".sh", ".log",
        ))
        or ctype.startswith("text/")
        or ctype in ("application/json", "application/xml")
    ):
        r = _extract_plaintext(contents)
        r["source_type"] = "text"
        return _finalize(r)

    try:
        r = _extract_plaintext(contents)
        r["source_type"] = "text"
        return _finalize(r)
    except Exception:
        raise ValueError(f"Unsupported file type: {content_type} ({filename})")


def _finalize(result: Dict[str, Any]) -> Dict[str, Any]:
    text = result.get("text", "") or ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.strip()

    result["text"] = text
    result["char_count"] = len(text)
    result.setdefault("pages", 1)
    result.setdefault("method", "unknown")
    return result